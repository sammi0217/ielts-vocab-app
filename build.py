#!/usr/bin/env python3
"""Inject vocabulary data from the Excel/Sheet export into app_template.html.

Usage:  python3 build.py [path/to/雅思單字進步系統.xlsx] [out.html]
Default: data/雅思單字進步系統.xlsx -> dist/index.html
Requires: pip install openpyxl
"""
import json, sys, datetime, os
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else "data/雅思單字進步系統.xlsx"
out = sys.argv[2] if len(sys.argv) > 2 else "dist/index.html"
FAM = {"未學習": 0, "學習中": 1, "熟悉": 2, "已掌握": 3}

wb = load_workbook(src)
ws = wb["單字庫"]
words = []
for r in range(2, ws.max_row + 1):
    v = [ws.cell(row=r, column=c).value for c in range(1, 15)]
    if not v[2]:
        continue
    d = v[11]
    d = d.strftime("%Y-%m-%d") if isinstance(d, (datetime.date, datetime.datetime)) else (d or "")
    words.append({
        "cat": v[1] or "", "w": v[2], "pos": v[3] or "", "zh": v[4] or "", "en": v[5] or "",
        "ex": v[6] or "", "exzh": v[7] or "", "syn": v[8] or "", "ant": v[9] or "-",
        "fam": FAM.get(v[10], 0), "date": d, "cnt": int(v[12] or 0), "note": v[13] or "",
    })

# carry over existing 7份輪替 log rows (複習紀錄 sheet, rows 14+, columns A..F)
log = []
if "複習紀錄" in wb.sheetnames:
    w2 = wb["複習紀錄"]
    for r in range(14, 60):
        row = [w2.cell(row=r, column=c).value for c in range(1, 7)]
        if any(x not in (None, "") for x in row[1:]):
            row = [x.strftime("%Y-%m-%d") if isinstance(x, (datetime.date, datetime.datetime)) else x for x in row]
            log.append(row)

meta = {"version": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), "total": len(words)}
html = open("app_template.html", encoding="utf-8").read()
j = lambda o: json.dumps(o, ensure_ascii=False).replace("</", "<\\/")
html = html.replace("__META__", j(meta)).replace("__WORDS__", j(words)).replace("__LOG__", j(log))
os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
open(out, "w", encoding="utf-8").write(html)
print(f"built {out}: {len(words)} words, {len(log)} log rows, version {meta['version']}")
